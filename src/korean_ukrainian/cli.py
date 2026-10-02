from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from .pipeline import analyze_korean, transliterate_korean


def _result(text: str, mode: str) -> dict[str, Any]:
    if mode == "decompose":
        return {"input": text, "analysis": analyze_korean(text)}
    if mode == "transliterate":
        return transliterate_korean(text)
    raise ValueError(f"unsupported mode: {mode}")


def _read_rows(path: Path, input_format: str, column: str, sheet: str | None) -> tuple[list[dict[str, Any]], list[str]]:
    if input_format == "txt":
        return [{"text": line.rstrip("\r\n")} for line in path.read_text(encoding="utf-8").splitlines() if line.strip()], ["text"]
    if input_format == "csv":
        import csv
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            rows = list(csv.DictReader(fh))
        if not rows or column not in rows[0]:
            raise ValueError(f"CSV column not found: {column!r}")
        return rows, list(rows[0])
    if input_format in {"xlsx", "xls"}:
        if input_format == "xls":
            raise ValueError("legacy .xls is not supported; convert to .xlsx")
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        values = list(ws.values)
        if not values:
            return [], []
        headers = [str(x) if x is not None else "" for x in values[0]]
        if column not in headers:
            raise ValueError(f"Excel column not found: {column!r}")
        return [dict(zip(headers, row)) for row in values[1:]], headers
    raise ValueError(f"unsupported input format: {input_format}")


def _write_rows(path: Path, output_format: str, rows: list[dict[str, Any]], headers: list[str]) -> None:
    if output_format == "jsonl":
        path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows) + ("\n" if rows else ""), encoding="utf-8")
        return
    if output_format == "json":
        path.write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding="utf-8")
        return
    if output_format == "csv":
        import csv
        with path.open("w", encoding="utf-8-sig", newline="") as fh:
            writer = csv.DictWriter(fh, fieldnames=headers, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        return
    if output_format == "xlsx":
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "results"
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h) for h in headers])
        wb.save(path)
        return
    if output_format == "txt":
        path.write_text("\n".join(str(r.get("ukrainian_orthography", "")) for r in rows) + ("\n" if rows else ""), encoding="utf-8")
        return
    raise ValueError(f"unsupported output format: {output_format}")


def run_batch(input_path: Path, output_path: Path, input_format: str, output_format: str, column: str, sheet: str | None, mode: str) -> int:
    rows, headers = _read_rows(input_path, input_format, column, sheet)
    output_rows: list[dict[str, Any]] = []
    extra = ["input", "ukrainian_orthography", "ipa", "analysis_status", "result_json"]
    for row in rows:
        text = str(row.get(column, "") or "").strip()
        if not text:
            result = {"input": text, "error": "empty input"}
        else:
            try:
                result = _result(text, mode)
            except Exception as exc:
                result = {"input": text, "error": f"{type(exc).__name__}: {exc}"}
        enriched = dict(row)
        enriched["input"] = text
        enriched["ukrainian_orthography"] = result.get("ukrainian_orthography", "")
        enriched["ipa"] = result.get("ipa", {}).get("ipa", "") if isinstance(result.get("ipa"), dict) else ""
        enriched["analysis_status"] = result.get("analysis_status", result.get("selection_status", ""))
        enriched["result_json"] = json.dumps(result, ensure_ascii=False)
        output_rows.append(enriched)
    if output_format == "txt":
        _write_rows(output_path, output_format, output_rows, headers + extra)
    else:
        _write_rows(output_path, output_format, output_rows, headers + [x for x in extra if x not in headers])
    return len(output_rows)


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="korean-ua", description="Korean → Ukrainian phonetic-grahemic analysis system")
    sub = p.add_subparsers(dest="command", required=True)

    one = sub.add_parser("analyze", help="Analyze one Korean string")
    one.add_argument("text")
    one.add_argument("--mode", choices=["decompose", "transliterate"], default="transliterate")

    batch = sub.add_parser("batch", help="Process TXT, CSV or XLSX in batch")
    batch.add_argument("input", type=Path)
    batch.add_argument("-o", "--output", type=Path, required=True)
    batch.add_argument("--input-format", choices=["txt", "csv", "xlsx", "xls"], default=None)
    batch.add_argument("--output-format", choices=["jsonl", "json", "csv", "xlsx", "txt"], default=None)
    batch.add_argument("--column", default="text", help="Input text column for CSV/XLSX")
    batch.add_argument("--sheet", default=None, help="Excel worksheet name")
    batch.add_argument("--mode", choices=["decompose", "transliterate"], default="transliterate")
    return p


def main() -> int:
    args = build_parser().parse_args()
    if args.command == "analyze":
        print(json.dumps(_result(args.text, args.mode), ensure_ascii=False, indent=2))
        return 0
    input_format = args.input_format or args.input.suffix.lower().lstrip(".")
    output_format = args.output_format or ("jsonl" if args.output.suffix.lower() in {".jsonl", ".ndjson"} else args.output.suffix.lower().lstrip("."))
    count = run_batch(args.input, args.output, input_format, output_format, args.column, args.sheet, args.mode)
    print(f"processed={count} output={args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
